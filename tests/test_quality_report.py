from openpyxl import load_workbook
from recon.report import create_report

def test_quality_report_has_independent_sheet(tmp_path):
    (tmp_path / "left.csv").write_text("id,amount\na,1\nb,\n", encoding="utf-8")
    (tmp_path / "right.csv").write_text("id,amount\na,1\nb,2\n", encoding="utf-8")
    rules = tmp_path / "rules.yaml"
    rules.write_text("left: left.csv\nright: right.csv\nkey: id\ncolumns: [amount]\nquality:\n  threshold: 95\n", encoding="utf-8")
    output = tmp_path / "report.xlsx"
    create_report(rules, output)
    book = load_workbook(output)
    assert "Quality" in book.sheetnames
    values = list(book["Quality"].values)
    assert values[0][1] == "score" and values[1][3] is False
