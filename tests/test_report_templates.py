import pytest
from openpyxl import load_workbook

from recon.report import create_report


def _rules(tmp_path):
    (tmp_path / "l.csv").write_text("id,v\na,1\n", encoding="utf-8")
    (tmp_path / "r.csv").write_text("id,v\na,2\n", encoding="utf-8")
    path = tmp_path / "rules.yaml"
    path.write_text("left: l.csv\nright: r.csv\nkey: id\ncolumns: [v]\n", encoding="utf-8")
    return path


@pytest.mark.parametrize("template", ["simple", "detailed", "management"])
def test_report_templates(tmp_path, template):
    out = tmp_path / f"{template}.xlsx"
    create_report(_rules(tmp_path), out, template=template)
    sheets = load_workbook(out).sheetnames
    assert "Differences" in sheets
    assert (template == "simple" and sheets == ["Summary", "Differences"]) or template != "simple"


def test_report_template_invalid(tmp_path):
    with pytest.raises(ValueError):
        create_report(_rules(tmp_path), tmp_path / "x.xlsx", template="bad")
