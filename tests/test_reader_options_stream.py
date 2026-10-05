import pytest
from openpyxl import Workbook

from recon.readers import ReaderOptions, iter_rows


def test_reader_options_validates_values():
    with pytest.raises(ValueError):
        ReaderOptions(header_row=-1)
    with pytest.raises(ValueError):
        ReaderOptions(max_rows=-1)
    with pytest.raises(ValueError):
        ReaderOptions(delimiter=",,")


def test_csv_stream_skips_header_blank_and_comments(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("# comment\nid,value\n\nA,1\nB,2\n", encoding="utf-8")
    rows = list(iter_rows(path, ReaderOptions(comment="#")))
    assert rows == [["A", "1"], ["B", "2"]]


def test_tsv_stream_respects_limit_and_header_offset(tmp_path):
    path = tmp_path / "data.tsv"
    path.write_text("说明\nid\tvalue\nA\t1\nB\t2\n", encoding="utf-8")
    rows = list(iter_rows(path, ReaderOptions(header_row=1, max_rows=1)))
    assert rows == [["A", "1"]]


def test_stream_rejects_unsupported_format(tmp_path):
    path = tmp_path / "data.json"
    path.write_text("[]", encoding="utf-8")
    with pytest.raises(Exception, match="流式读取"):
        list(iter_rows(path))


def test_xlsx_stream_uses_read_only_rows(tmp_path):
    path = tmp_path / "data.xlsx"
    book = Workbook()
    book.active.append(["id", "value"])
    book.active.append(["A", 1])
    book.active.append(["B", 2])
    book.save(path)
    assert list(iter_rows(path, ReaderOptions(max_rows=1))) == [["A", 1]]
