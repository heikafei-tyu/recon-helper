from openpyxl import load_workbook

from recon.report import create_report


def test_report_and_incremental_skip(tmp_path):
    (tmp_path / "custom-left.csv").write_text("id,amount\na,1\n", encoding="utf-8")
    (tmp_path / "custom-right.csv").write_text("id,amount\na,2\n", encoding="utf-8")
    rules = tmp_path / "rules.yaml"
    rules.write_text("left: custom-left.csv\nright: custom-right.csv\nkey: id\ncolumns: [amount]\n", encoding="utf-8")
    output = tmp_path / "report.xlsx"
    assert create_report(rules, output)["skipped"] is False
    book = load_workbook(output)
    assert book.sheetnames == ["Differences", "Summary"]
    assert book["Differences"].freeze_panes == "A2"
    assert book["Summary"]["A1"].value == "metric"
    assert book["Summary"]["A2"].value == "tool_version"
    metrics = {row[0].value: row[1].value for row in book["Summary"].iter_rows(min_row=2)}
    assert metrics["run_at_utc"]
    assert len(metrics["left_sha256"]) == 64
    assert create_report(rules, output, incremental=True)["skipped"] is True

def test_report_refuses_overwrite_without_force(tmp_path):
    (tmp_path / "orders.csv").write_text("id,amount\na,1\n", encoding="utf-8")
    (tmp_path / "bank.csv").write_text("id,amount\na,2\n", encoding="utf-8")
    rules = tmp_path / "rules.yaml"
    rules.write_text("left: orders.csv\nright: bank.csv\nkey: id\ncolumns: [amount]\n", encoding="utf-8")
    output = tmp_path / "report.xlsx"
    create_report(rules, output)
    import pytest
    with pytest.raises(ValueError, match="--force"):
        create_report(rules, output)
