import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill

from .engine import run_rules
from .fingerprint_store import FingerprintStore
from .quality import assess_file


def create_html_report(rules_file, output):
    result = run_rules(rules_file)
    for difference in result["differences"]:
        difference.setdefault("review_status", "未处理")
        difference.setdefault("review_note", "")
    rows = result["differences"]
    columns = ["key", "status", "review_status", "review_note", "column", "left_value", "right_value", "difference"]
    def esc(value):
        return str(value).replace("&", "&amp;").replace("<", "&lt;").replace('"', "&quot;")
    body = "".join("<tr data-status='{}' data-table='{}' title='上下文：{}'>".format(esc(row.get("status", "")), esc(row.get("left_table", row.get("from", ""))), esc(row.get("key", ""))) + "".join(f"<td>{esc(row.get(column, ''))}</td>" for column in columns) + "</tr>" for row in rows)
    header = "".join(f"<th data-column='{column}'>{column} ↕</th>" for column in columns)
    html = f"""<!doctype html><meta charset='utf-8'><title>recon-helper report</title>
<style>body{{font:14px system-ui;margin:2rem}}table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #ddd;padding:6px}}th{{cursor:pointer;background:#eef2f7}}tr[data-status='mismatch']{{background:#fff2cc}}.controls{{display:flex;gap:1rem;margin:1rem 0}}label{{font-weight:600}}</style>
<h1>Recon report</h1><div class='controls'><label>表 <select id='table'><option value=''>全部</option></select></label><label>严重程度 <select id='status'><option value=''>全部</option><option>mismatch</option><option>within_tolerance</option><option>left_only</option><option>right_only</option></select></label></div>
<table id='report'><thead><tr>{header}</tr></thead><tbody>{body}</tbody></table>
<script>(function(){{const table=document.querySelector('#report'), rows=[...table.tBodies[0].rows], tableSelect=document.querySelector('#table');[...new Set(rows.map(r=>r.dataset.table).filter(Boolean))].forEach(v=>tableSelect.add(new Option(v,v)));function filter(){{const t=tableSelect.value,s=document.querySelector('#status').value;rows.forEach(r=>r.hidden=(t&&r.dataset.table!==t)||(s&&r.dataset.status!==s));}}tableSelect.onchange=filter;document.querySelector('#status').onchange=filter;table.tHead.addEventListener('click',e=>{{const cell=e.target.closest('th');if(!cell)return;const i=cell.cellIndex, asc=cell.dataset.asc!=='1';rows.sort((a,b)=>String(a.cells[i]?.textContent).localeCompare(String(b.cells[i]?.textContent),undefined,{{numeric:true}})*(asc?1:-1));rows.forEach(r=>table.tBodies[0].appendChild(r));cell.dataset.asc=asc?'1':'0';}});}})();</script>"""
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(html, encoding="utf-8")
    return {"output": str(output), "differences": len(rows)}


def _digest(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()

def create_report(rules_file, output, incremental=False, force=False, template="detailed"):
    if template not in {"simple", "detailed", "management"}:
        raise ValueError("未知报告模板：simple、detailed、management")
    rules_path = Path(rules_file)
    output = Path(output)
    import yaml
    config = yaml.safe_load(rules_path.read_text(encoding="utf-8-sig"))
    source_names = [config.get("left"), config.get("right")]
    for check in config.get("checks", []):
        source_names.extend([check.get("detail"), check.get("summary"), check.get("file")])
        source_names.extend(check.get("tables", []))
    sources = [rules_path, *[rules_path.parent / name for name in source_names if name]]
    sources = list(dict.fromkeys(sources))
    manifest = {str(p): _digest(p) for p in sources if p.exists()}
    marker = output.with_suffix(output.suffix + ".manifest.json")
    if output.exists() and not incremental and not force:
        raise ValueError(f"报告已存在：{output}；如需覆盖请使用 --force")
    db_path = output.parent / "recon_history.db"
    with FingerprintStore(db_path) as store:
        if marker.exists() and not store.get(str(rules_path)):
            store.migrate_json(marker)
        cached = all(store.get(path) == digest for path, digest in manifest.items())
    if incremental and output.exists() and cached:
        try:
            existing = load_workbook(output, read_only=True)
            valid = {"Differences", "Summary"}.issubset(existing.sheetnames)
            existing.close()
        except Exception:
            valid = False
        if valid:
            return {"output": str(output), "skipped": True}
    result = run_rules(rules_file)
    for difference in result["differences"]:
        difference.setdefault("review_status", "未处理")
        difference.setdefault("review_note", "")
    quality_scores = []
    quality_config = config.get("quality", {}) if isinstance(config.get("quality", {}), dict) else {}
    threshold = quality_config.get("threshold", 70)
    key_columns = config.get("keys") or ([config.get("key")] if config.get("key") else None)
    for source_name in source_names[:2]:
        if source_name and (rules_path.parent / source_name).exists():
            quality_scores.append(assess_file(rules_path.parent / source_name, key_columns, threshold=threshold))
    result["quality_scores"] = quality_scores
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    rows = result["differences"]
    preferred = ["key", "status", "review_status", "review_note", "left_table", "right_table", "left_row", "right_row", "column", "left_value", "right_value", "difference", "tolerance", "tolerance_type", "tolerance_rule", "rule_priority"]
    columns = [key for key in preferred if any(key in row for row in rows)] or ["status"]
    def add_detail_sheet(title, selected):
        sheet = book.create_sheet(title)
        sheet.append(columns)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{chr(64 + min(len(columns), 26))}{len(selected) + 1}"
        fills = PatternFill("solid", fgColor="FFF2CC")
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in selected:
            sheet.append([" / ".join(row[k]) if isinstance(row.get(k), (tuple, list)) else row.get(k) for k in columns])
            if row.get("status") in ("mismatch", "total_mismatch", "chain_mismatch", "data_missing", "left_only", "right_only"):
                for cell in sheet[sheet.max_row]:
                    cell.fill = fills
        for column_cells in sheet.columns:
            letter = column_cells[0].column_letter
            sheet.column_dimensions[letter].width = min(40, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))
        return sheet

    add_detail_sheet("Differences", rows)
    detail_names = []
    for name in source_names if template != "simple" else []:
        if not name:
            continue
        title = Path(name).stem[:25] or "table"
        if title in {"Summary", "Differences"} or title in detail_names:
            title = f"{title}_{len(detail_names) + 1}"
        detail_names.append(title)
        selected = [row for row in rows if Path(str(row.get("left_table", ""))).stem == Path(name).stem or Path(str(row.get("right_table", ""))).stem == Path(name).stem or row.get("from") == name or row.get("to") == name]
        add_detail_sheet(title, selected)
    summary.append(["metric", "value"])
    counts = {status: sum(r.get("status") == status for r in rows) for status in ("mismatch", "within_tolerance", "left_only", "right_only")}
    summary_rows = [("tool_version", "0.1.0"), ("run_at_utc", datetime.now(timezone.utc).isoformat()), ("rules_file", str(rules_path)), ("left_file", str(sources[1]) if len(sources) > 1 else None), ("right_file", str(sources[2]) if len(sources) > 2 else None), ("left_sha256", manifest.get(str(sources[1])) if len(sources) > 1 else None), ("right_sha256", manifest.get(str(sources[2])) if len(sources) > 2 else None), ("left_rows", result.get("left_rows", 0)), ("right_rows", result.get("right_rows", 0)), ("differences", len(rows)), ("conclusion", "通过" if not rows else "存在差异"), *counts.items()]
    for key, value in summary_rows:
        summary.append([key, value])
    if quality_scores:
        summary.append(["quality_scores", ""])
        for quality in quality_scores:
            row = summary.max_row + 1
            summary.append([quality["source"] or "input", quality["score"]])
            if not quality["passed"]:
                for cell in summary[row]:
                    cell.fill = PatternFill("solid", fgColor="FFC7CE")
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        chart_path = output.with_suffix(".summary.png")
        labels = list(counts); values = list(counts.values())
        figure, axis = plt.subplots(figsize=(5, 3)); axis.bar(labels, values, color="#4C78A8"); axis.set_title("差异分布"); figure.tight_layout(); figure.savefig(chart_path); plt.close(figure)
        if template != "simple":
            summary.add_image(XLImage(str(chart_path)), "D2")
    except ImportError:
        pass
    if template == "management":
        summary.append(["建议", "优先处理缺失数据与超容差差异；复核输入指纹后再提交。"])
    summary.freeze_panes = "A2"
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    marker.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    with FingerprintStore(db_path) as store:
        store.put_many(manifest)
    return {"output": str(output), "skipped": False, "differences": len(rows)}
