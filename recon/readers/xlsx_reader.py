from openpyxl import load_workbook

from ..errors import ReadError
from ..model import Table


def read(path, encoding=None, sheet_name=None, sheet_index=0, skip_rows=0):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name is not None:
            if sheet_name not in book.sheetnames:
                raise ReadError(f"工作表不存在：{sheet_name}")
            sheet = book[sheet_name]
        else:
            try:
                sheet = book.worksheets[sheet_index]
            except IndexError as exc:
                raise ReadError(f"工作表序号不存在：{sheet_index}") from exc
        def records():
            for cells in sheet.iter_rows():
                if any(cell.data_type == "f" for cell in cells):
                    raise ReadError("XLSX 包含公式，请先转为已核验的值")
                yield [cell.value for cell in cells]
        rows = records()
        for _ in range(skip_rows):
            next(rows, None)
        header = next(rows, None)
        if header is None:
            raise ReadError("XLSX 工作表为空")
        table = Table.from_records(header, rows)
        return Table(table.columns, table.rows, str(path), "xlsx", None, None, sheet.title)
    finally:
        book.close()


def read_sheets(path):
    """读取工作簿中所有非空工作表，返回 sheet 名到 Table 的映射。"""
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        result = {}
        for sheet in book.worksheets:
            def records():
                for cells in sheet.iter_rows():
                    if any(cell.data_type == "f" for cell in cells):
                        raise ReadError("XLSX 包含公式，请先转为已核验的值")
                    yield [cell.value for cell in cells]
            rows = records()
            header = next(rows, None)
            if header is not None:
                table = Table.from_records(header, rows)
                result[sheet.title] = Table(table.columns, table.rows, str(path), "xlsx", None, None, sheet.title)
        if not result:
            raise ReadError("XLSX 工作簿没有非空工作表")
        return result
    finally:
        book.close()
