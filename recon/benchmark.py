import csv
import tempfile
import time
import tracemalloc
from pathlib import Path


def stream_csv(path):
    with Path(path).open("r", encoding="utf-8-sig", newline="") as stream:
        yield from csv.DictReader(stream)


def stream_xlsx(path, sheet_name=None):
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[sheet_name] if sheet_name else book.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return
        for values in rows:
            yield dict(zip(header, values))
    finally:
        book.close()


def benchmark(path, key="id", check_duplicates=True, timeout=None, progress=False):
    if timeout is not None and timeout < 0:
        raise ValueError("timeout 不能为负数")
    path = Path(path)
    started = time.perf_counter()
    tracemalloc.start()
    rows = 0
    seen = set()
    duplicates = 0
    total_rows = None
    if progress:
        with path.open("rb") as stream:
            total_rows = max(1, sum(1 for _ in stream) - 1)
    iterator = stream_xlsx(path) if path.suffix.lower() == ".xlsx" else stream_csv(path)
    for row in iterator:
        rows += 1
        if progress and rows % 10000 == 0:
            percent = min(100.0, rows / total_rows * 100)
            print(f"processed_rows={rows} progress={percent:.1f}%")
        value = row.get(key)
        if check_duplicates and value in seen:
            duplicates += 1
        if check_duplicates:
            seen.add(value)
        if timeout is not None and time.perf_counter() - started > timeout:
            tracemalloc.stop()
            raise TimeoutError(f"超过 {timeout} 秒")
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    return {
        "file": str(path),
        "format": path.suffix.lower().lstrip("."),
        "bytes": path.stat().st_size,
        "rows": rows,
        "duplicates": duplicates if check_duplicates else None,
        "seconds": round(time.perf_counter() - started, 6),
        "peak_memory_mb": round(peak / 1024 / 1024, 3),
        "duplicate_check": check_duplicates,
    }


def generate_benchmark_csv(path, rows=100_000):
    path = Path(path)
    with path.open("w", encoding="utf-8", newline="") as stream:
        stream.write("id,amount\n")
        for index in range(rows):
            stream.write(f"{index},{index * 1.25:.2f}\n")
    return path


def benchmark_generated(rows=100_000, timeout=None, progress=False):
    with tempfile.TemporaryDirectory(prefix="recon-bench-") as directory:
        path = generate_benchmark_csv(Path(directory) / "generated.csv", rows)
        optimized = benchmark(path, timeout=timeout, progress=progress)
        baseline = benchmark_baseline(path)
        return {
            "rows_requested": rows,
            "optimized": optimized,
            "baseline": baseline,
            "comparison": {"seconds_saved": round(baseline["seconds"] - optimized["seconds"], 6), "memory_saved_mb": round(baseline["peak_memory_mb"] - optimized["peak_memory_mb"], 3)},
        }


def benchmark_baseline(path):
    """真实执行一次全量载入，作为流式方案的可重复对照。"""
    started = time.perf_counter()
    tracemalloc.start()
    if Path(path).suffix.lower() == ".xlsx":
        rows = list(stream_xlsx(path))
    else:
        rows = list(stream_csv(path))
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    return {"method": "full-table", "rows": len(rows), "seconds": round(time.perf_counter() - started, 6), "peak_memory_mb": round(peak / 1024 / 1024, 3)}


def memory_curve(path, interval=0.1):
    """用 memory_profiler 采样流式读取过程的内存曲线。"""
    try:
        from memory_profiler import memory_usage
    except ImportError as exc:
        raise RuntimeError("请安装 memory-profiler 后采样内存曲线") from exc
    samples = memory_usage(
        (
            benchmark,
            (path,),
            {},
        ),
        interval=interval,
    )
    return {
        "samples_mb": [round(value, 3) for value in samples],
        "peak_mb": round(max(samples), 3),
        "minimum_mb": round(min(samples), 3),
    }
